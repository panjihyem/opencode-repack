"""HYPER_AI_MANAGED: public/synthetic report deck from one data file; no network."""
import json
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt


def build(source: Path, output: Path):
    if output.exists() or output.resolve() == source.resolve():
        raise ValueError("Output already exists; source must remain untouched")
    data = json.loads(source.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or data.get("data_class") not in ("public", "synthetic"):
        raise ValueError("Only public/synthetic data accepted")
    title, caption, rows = data.get("title"), data.get("source"), data.get("rows")
    if not isinstance(title, str) or not title.strip() or not isinstance(caption, str) or not caption.strip() or not isinstance(rows, list) or not rows:
        raise ValueError("Title, source and nonempty rows required")
    if any(not isinstance(row, dict) or not isinstance(row.get("label"), str) or not isinstance(row.get("value"), (int, float)) or isinstance(row.get("value"), bool) for row in rows):
        raise ValueError("Each row requires label and numeric value")
    if len(rows) > 12 or len(title) > 95 or len(caption) > 180:
        raise ValueError("Deck exceeds bounded layout")

    output.parent.mkdir(parents=True, exist_ok=True)
    chart_file = output.with_suffix(".chart.png")
    if chart_file.exists():
        raise ValueError("Chart image already exists")
    labels = [r["label"] for r in rows]
    values = [r["value"] for r in rows]
    figure, axis = plt.subplots(figsize=(10, 4))
    axis.bar(labels, values, color="#2563eb")
    axis.set_title(title)
    axis.tick_params(axis="x", labelrotation=30)
    figure.tight_layout()
    figure.savefig(chart_file, dpi=150)
    plt.close(figure)

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = title
    first = deck.slides.add_slide(deck.slide_layouts[6])
    box = first.shapes.add_textbox(Inches(.7), Inches(.45), Inches(11.9), Inches(.8))
    box.text_frame.text = title
    box.text_frame.paragraphs[0].font.size = Pt(30)
    first.shapes.add_picture(str(chart_file), Inches(.85), Inches(1.5), width=Inches(11.5), height=Inches(4.6))
    first.shapes.add_textbox(Inches(.85), Inches(6.4), Inches(11.5), Inches(.55)).text = f"Sumber: {caption}"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(.7), Inches(.45), Inches(11.9), Inches(.8)).text = title
    series = CategoryChartData()
    series.categories = labels
    series.add_series("Nilai", values)
    second.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.85), Inches(1.4), Inches(11.5), Inches(4.8), series)
    second.shapes.add_textbox(Inches(.85), Inches(6.4), Inches(11.5), Inches(.55)).text = f"Sumber: {caption}"
    deck.save(output)
    return output, chart_file


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: deck.py <public-data.json> <new-output.pptx>")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
